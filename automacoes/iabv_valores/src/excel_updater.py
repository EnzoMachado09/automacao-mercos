from __future__ import annotations

import logging
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from .name_matcher import (
    ANO_PARA_COLUNA,
    ResultadoMatch,
    construir_lookup,
    encontrar_match,
)
from .pdf_parser import PedidoInfo, parsear_pasta

LOGGER = logging.getLogger(__name__)


@dataclass
class Config:
    pdfs_dir: Path
    excel_path: Path
    coluna_ano: str          # "2024", "2025" ou "2026"
    limiar_fuzzy: float = 0.75
    dry_run: bool = False
    blacklist_pdf: set[str] = None  # nomes normalizados do PDF que devem ser ignorados

    def __post_init__(self) -> None:
        if self.coluna_ano not in ANO_PARA_COLUNA:
            raise ValueError(f"coluna_ano deve ser '2024', '2025' ou '2026', recebido: {self.coluna_ano!r}")
        if self.blacklist_pdf is None:
            self.blacklist_pdf = set()


@dataclass
class ResumoExecucao:
    total_pdfs: int = 0
    pdfs_com_erro: int = 0
    clientes_atualizados: int = 0
    matches_exatos: int = 0
    matches_alias: int = 0
    matches_substring: int = 0
    matches_fuzzy: int = 0
    nao_identificados: list[str] = None

    def __post_init__(self) -> None:
        if self.nao_identificados is None:
            self.nao_identificados = []


def executar(config: Config) -> ResumoExecucao:
    resumo = ResumoExecucao()

    # --- 1. Parsear PDFs ---
    LOGGER.info(f"Lendo PDFs em: {config.pdfs_dir}")
    pedidos, erros = parsear_pasta(config.pdfs_dir)
    resumo.total_pdfs = len(pedidos) + len(erros)
    resumo.pdfs_com_erro = len(erros)

    if erros:
        for e in erros:
            LOGGER.warning(f"  Erro ao ler PDF: {e.name}")

    if not pedidos:
        LOGGER.warning("Nenhum PDF válido encontrado.")
        return resumo

    LOGGER.info(f"{len(pedidos)} PDFs lidos com sucesso.")

    # --- 2. Carregar Excel ---
    LOGGER.info(f"Carregando Excel: {config.excel_path}")
    wb = openpyxl.load_workbook(config.excel_path)
    ws = wb.active
    entradas = construir_lookup(ws)
    LOGGER.info(f"{len(entradas)} clientes carregados da planilha.")

    coluna_idx = ANO_PARA_COLUNA[config.coluna_ano]

    # --- 3. Matching e acumulação ---
    # totais[linha_excel] = soma dos valores dos pedidos que bateram
    totais: dict[int, float] = defaultdict(float)
    # detalhes para log
    detalhes_match: list[tuple[PedidoInfo, ResultadoMatch]] = []

    for pedido in pedidos:
        if pedido.nome_cliente in config.blacklist_pdf:
            LOGGER.info(f"  [blacklist] {pedido.nome_cliente_raw!r} ignorado (PDF: {pedido.pdf_path.name})")
            continue

        match = encontrar_match(pedido.nome_cliente, entradas, config.limiar_fuzzy)
        if match is None:
            LOGGER.warning(f"  NAO IDENTIFICADO: {pedido.nome_cliente_raw!r} (PDF: {pedido.pdf_path.name})")
            resumo.nao_identificados.append(f"{pedido.nome_cliente_raw} | {pedido.pdf_path.name} | R$ {pedido.valor:,.2f}")
            continue

        totais[match.entrada.linha] += pedido.valor
        detalhes_match.append((pedido, match))

        tipo_icon = {"exato": "OK", "alias": "OK", "substring": "~", "fuzzy": "(!)"}.get(match.tipo, "?")
        LOGGER.info(
            f"  {tipo_icon} [{match.tipo}] {pedido.nome_cliente_raw!r}"
            f" -> {match.entrada.nome_original!r}"
            f" | R$ {pedido.valor:,.2f}"
            f" (score={match.score:.2f})"
        )

        if match.tipo == "exato" or match.tipo == "alias":
            resumo.matches_exatos += 1 if match.tipo == "exato" else 0
            resumo.matches_alias += 1 if match.tipo == "alias" else 0
        elif match.tipo == "substring":
            resumo.matches_substring += 1
        else:
            resumo.matches_fuzzy += 1

    # --- 4. Aplicar no Excel ---
    resumo.clientes_atualizados = len(totais)

    if config.dry_run:
        LOGGER.info("[DRY RUN] Valores que seriam somados:")
        for linha, total in sorted(totais.items()):
            nome = ws.cell(row=linha, column=1).value
            atual = ws.cell(row=linha, column=coluna_idx).value or 0
            LOGGER.info(f"  Linha {linha}: {nome!r} | atual={atual:,.2f} | +{total:,.2f} = {atual+total:,.2f}")
        LOGGER.info("[DRY RUN] Excel NÃO foi salvo.")
    else:
        for linha, total in totais.items():
            cell = ws.cell(row=linha, column=coluna_idx)
            atual = cell.value if isinstance(cell.value, (int, float)) else 0
            cell.value = round(atual + total, 2)

        wb.save(config.excel_path)
        LOGGER.info(f"Excel salvo: {config.excel_path}")

    return resumo
